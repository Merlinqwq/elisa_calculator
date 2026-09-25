"""Import labeled subsets of a 96-well plate without inferring assay roles."""

from __future__ import annotations

from io import BytesIO
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROWS = "ABCDEFGH"
OVERFLOW_TOKENS = {"OVRFLW", "OVERFLOW", "OVER RANGE", "OVER-RANGE", "SATURATED"}


class ImportErrorDetail(ValueError):
    pass


def _header_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and isfinite(value) and int(value) == value:
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _candidates(sheet):
    found = []
    for header_row in range(1, min(sheet.max_row, 60) + 1):
        for first_col in range(2, min(sheet.max_column, 40) + 1):
            first_number = _header_number(sheet.cell(header_row, first_col).value)
            if first_number is None or not 1 <= first_number <= 12:
                continue
            if any(header_row in earlier_rows.values() and first_col in earlier_cols
                   for _, _, earlier_cols, earlier_rows in found):
                continue  # A reading of exactly 1, 2, etc. is not another header.
            previous = _header_number(sheet.cell(header_row, first_col - 1).value) if first_col > 2 else None
            if previous is not None and 1 <= previous < first_number <= 12:
                continue  # Keep the maximal numbered run, not each suffix of it.
            label_col = first_col - 1
            row_index = {}
            duplicated = False
            for r in range(header_row + 1, min(sheet.max_row, header_row + 20) + 1):
                label = sheet.cell(r, label_col).value
                if isinstance(label, str) and label.strip().upper() in ROWS:
                    label = label.strip().upper()
                    if label in row_index:
                        duplicated = True
                        break
                    row_index[label] = r
            if duplicated or not row_index:
                continue
            if list(row_index) != sorted(row_index, key=ROWS.index):
                continue
            cols = []
            last_number = 0
            malformed = False
            for c in range(first_col, min(sheet.max_column, first_col + 23) + 1):
                value = sheet.cell(header_row, c).value
                if value is None or isinstance(value, str) and not value.strip():
                    continue  # Spacer columns do not change physical well numbers.
                number = _header_number(value)
                if number is not None and 1 <= number <= 12 and number <= last_number:
                    malformed = True
                    break
                if number is None or not 1 <= number <= 12:
                    break
                cols.append(c)
                last_number = number
            if cols and not malformed:
                found.append((header_row, first_col, cols, row_index))
    return found


def _measurement(value):
    if value is None or isinstance(value, str) and not value.strip():
        return "empty", None
    if isinstance(value, str) and value.strip().upper() in OVERFLOW_TOKENS:
        return "overflow", None
    if not isinstance(value, bool):
        try:
            number = float(value)
            if isfinite(number):
                return "numeric", number
        except (TypeError, ValueError, OverflowError):
            pass
    return "invalid", None


def import_workbook(data: bytes, filename: str = "plate.xlsx", sheet_name: str | None = None) -> dict:
    if not filename.lower().endswith(".xlsx"):
        raise ImportErrorDetail("Upload an .xlsx workbook.")
    if len(data) > 20_000_000:
        raise ImportErrorDetail("Workbook exceeds the 20 MB local import limit.")
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ImportErrorDetail(f"Could not read Excel workbook: {exc}") from exc
    matches = []
    for sheet in workbook.worksheets:
        for match in _candidates(sheet):
            matches.append((sheet, *match))
    candidate_sheets = list(dict.fromkeys(s.title for s, *_ in matches))
    if sheet_name is None and len(candidate_sheets) > 1:
        workbook.close()
        return {
            "requires_sheet_selection": True,
            "filename": Path(filename).name,
            "candidate_sheets": candidate_sheets,
            "message": "More than one sheet looks like a plate. Choose the sheet with OD readings.",
        }
    if sheet_name is not None:
        matches = [match for match in matches if match[0].title == sheet_name]
    if len(matches) != 1:
        locations = [f"{s.title}!{get_column_letter(c)}{r}" for s, r, c, _, _ in matches]
        workbook.close()
        raise ImportErrorDetail(
            f"Expected one unambiguous labeled plate block{f' on {sheet_name}' if sheet_name else ''} "
            "with numbered well headers 1–12; found " + (", ".join(locations) if locations else "none") + "."
        )
    sheet, header_row, first_col, cols, row_index = matches[0]
    column_map = {_header_number(sheet.cell(header_row, c).value): c for c in cols}
    wells = []
    for row_letter in ROWS:
        for col_number in range(1, 13):
            source_cell = None
            raw = None
            if row_letter in row_index and col_number in column_map:
                cell = sheet.cell(row_index[row_letter], column_map[col_number])
                raw = cell.value
                source_cell = f"{get_column_letter(column_map[col_number])}{row_index[row_letter]}"
            status, raw_od = _measurement(raw)
            if raw is not None and not isinstance(raw, (str, int, float, bool)):
                raw = str(raw)
            wells.append({
                "well": f"{row_letter}{col_number}", "row": row_letter,
                "column": col_number, "raw_value": raw, "raw_od": raw_od,
                "measurement_status": status, "source_cell": source_cell,
            })
    plate_columns = sorted(column_map)
    plate_rows = list(row_index)
    warning = (f"Imported rows {', '.join(plate_rows)} and columns {', '.join(map(str, plate_columns))}; "
               "unprovided wells are empty in the plate view.")
    workbook.close()
    return {
        "filename": Path(filename).name, "sheet": sheet.title,
        "header_row": header_row, "imported_columns": len(cols),
        "plate_columns": plate_columns, "plate_rows": plate_rows,
        "wells": wells,
        "warnings": [warning] if len(cols) < 12 or len(plate_rows) < 8 else [],
    }
