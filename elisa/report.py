"""Scientific report workbook and labeled logistic curve plots."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .curve import four_pl


def _figure(curve: dict, standard_wells: list[dict]):
    fig, ax = plt.subplots(figsize=(7.5, 4.4), layout="constrained")
    fig.patch.set_facecolor("white")
    numeric = [w for w in standard_wells if w["include"] and w["corrected_od"] is not None]
    if numeric:
        ax.scatter([w["nominal_concentration"] for w in numeric], [w["corrected_od"] for w in numeric],
                   s=35, color="#2b6f9c", label="Individual standard wells", zorder=3)
    for level in curve.get("level_details", []):
        ax.errorbar(level["nominal_concentration"], level["mean_od"], yerr=level["sd_od"],
                    fmt="o", color="#d36d2b", capsize=3, markersize=5,
                    label="Mean ± SD" if level is curve["level_details"][0] else None, zorder=4)
    params = curve.get("parameters")
    if params and curve.get("range_low") and curve.get("range_high"):
        x = np.geomspace(curve["range_low"], curve["range_high"], 200)
        y = four_pl(x, **params)
        ax.plot(x, y, color="#164b6a", lw=2, label=f"Fitted {curve['model_type']} within tested range")
        ax.axvspan(curve["range_low"], curve["range_high"], color="#daf0eb", alpha=.35, zorder=0)
    overflow = [w for w in standard_wells if w["include"] and w["measurement_status"] == "overflow"]
    for w in overflow:
        ax.annotate(f"{w['well']} overflow", xy=(w["nominal_concentration"], .97),
                    xycoords=("data", "axes fraction"), rotation=90, va="top", ha="right",
                    fontsize=7, color="#ae3b3b")
    fit_label = "FIT OK" if curve["status"] == "valid" else "NO FIT"
    ax.set_title(f"Standard curve {curve['curve_id']} — {fit_label}; QC {curve['qc_status'].upper()}")
    ax.set_xlabel(f"Nominal concentration ({curve['units']})")
    ax.set_ylabel("Corrected OD")
    ax.grid(alpha=.2)
    if numeric:
        ax.legend(fontsize=8, loc="best")
    if curve.get("range_low"):
        ax.set_xscale("symlog", linthresh=curve["range_low"] / 10)
    return fig


def curve_plot_bytes(curve: dict, well_results: list[dict], fmt="svg") -> bytes:
    wells = [w for w in well_results if w["role"] == "standard" and w["curve_id"] == curve["curve_id"]]
    fig = _figure(curve, wells)
    output = BytesIO()
    fig.savefig(output, format=fmt, dpi=150)
    plt.close(fig)
    return output.getvalue()


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, len(rows)+1)}"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="174A67")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in rows:
        ws.append(["; ".join(v) if isinstance(v, list) else v for v in row])
    for col_index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col_index)].width = (widths[col_index - 1] if widths else min(38, max(13, len(header)+3)))
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, float):
                if math.isfinite(cell.value):
                    cell.number_format = "0.##########"
                else:
                    cell.value = None
        if cells[0].row % 2 == 0:
            for cell in cells:
                cell.fill = PatternFill("solid", fgColor="F1F6F8")


def workbook_bytes(result: dict, layout: dict) -> bytes:
    wb = Workbook()
    plate = wb.active
    plate.title = "Concentration Plate"
    plate.append(["Well / column"] + list(range(1, 13)))
    by_well = {w["well"]: w for w in result["well_results"]}
    for row_name in "ABCDEFGH":
        row = [row_name]
        for column in range(1, 13):
            well = by_well[f"{row_name}{column}"]
            if well["role"] == "blank":
                value = "BLANK"
            elif well["role"] == "standard":
                value = f"REF {well['nominal_concentration']:g}" + ("\nOVRFLW" if well["measurement_status"] == "overflow" else "")
            elif well["role"] == "unknown":
                value = well["well_adjusted_concentration"]
                if value is None:
                    if well["result_status"].startswith("below_"):
                        value = 0
                    elif well["measurement_status"] == "overflow":
                        value = "OVRFLW"
                    elif well["result_status"] == "invalid_curve":
                        value = "QC FAIL"
                    else:
                        value = "—"
            else:
                value = None
            row.append(value)
        plate.append(row)
    plate.append(["0 means below lowest tested standard, not measured zero. Unknown values include dilution factor."])
    plate.append(["Amber cells need QC review; see Unknown Results and Standard Curves."])
    plate.freeze_panes = "B2"
    plate.column_dimensions["A"].width = 82
    for column in range(2, 14):
        plate.column_dimensions[get_column_letter(column)].width = 17
    for cell in plate[1]:
        cell.fill = PatternFill("solid", fgColor="174A67")
        cell.font = Font(color="FFFFFF", bold=True)
    for cells in plate.iter_rows(min_row=2, max_row=9):
        for cell in cells[1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.##########"
    curve_by_id = {curve["curve_id"]: curve for curve in result["curves"]}
    for row_number, row_name in enumerate("ABCDEFGH", start=2):
        for column in range(1, 13):
            well = by_well[f"{row_name}{column}"]
            curve = curve_by_id.get(well["curve_id"])
            if well["role"] == "unknown" and (well["result_status"] != "ok" or
                                               (curve and curve["qc_status"] == "fail")):
                cell = plate.cell(row_number, column + 1)
                cell.fill = PatternFill("solid", fgColor="FFE3AD")
                cell.comment = Comment(
                    f"{well['well']}: {well['result_status']}; curve QC "
                    f"{curve['qc_status'] if curve else 'unavailable'}. See detailed sheets.",
                    "ELISA tool")
    summary = wb.create_sheet("Summary")
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    method_meaning = {"A": "average individual well concentrations", "B": "invert mean replicate OD"}[result["settings"]["primary_method"]]
    overview = [
        ["Source file", result["filename"]], ["Source sheet", result["sheet"]],
        ["Generated", generated], ["Software version", __version__],
        ["Blank policy", result["settings"]["blank_policy"]],
        ["Primary method", f"{result['settings']['primary_method']}: {method_meaning}"],
        ["Equations", "See Standard Curves for each fitted equation"],
        ["Range policy", "Interpolation within positive standard concentrations and usable response; no extrapolation"],
        ["Assay validation", "Software QC only; reported bounds are not validated LLOQ/ULOQ"],
    ]
    overview += [[f"Standard curve {c['curve_id']}", f"fit {c['status']}; QC {c['qc_status']}; {c['n_levels']} usable levels; {c['units']}; " + "; ".join(c["flags"])] for c in result["curves"]]
    overview += [["Global warning", value] for value in result["errors"]]
    _sheet(summary, ["Item", "Value"], overview, [25, 105])

    unknown = wb.create_sheet("Unknown Results")
    unknown_headers = ["Sample ID", "Replicate group", "Standard curve ID", "Units", "Dilution factor", "Assigned n", "Included n", "Valid n", "Wells", "Mean corrected OD", "OD SD", "OD CV %", "Method A measured (average well concentrations)", "Method B measured (invert mean OD)", "Primary method", "Primary measured", "Dilution-adjusted", "Partial valid mean (diagnostic)", "Valid well concentration SD", "Valid well concentration CV %", "Status", "Flags"]
    unknown_rows = [[r[k] for k in ("sample_id", "replicate_group", "curve_id", "units", "dilution_factor", "n_assigned", "n", "n_valid", "wells", "mean_corrected_od", "sd_corrected_od", "cv_corrected_od_pct", "method_a_measured_concentration", "method_b_measured_concentration", "primary_method", "primary_measured_concentration", "dilution_adjusted_concentration", "partial_valid_mean_concentration", "sd_concentration_method_a", "cv_concentration_method_a_pct", "status", "flags")] for r in result["unknown_results"]]
    _sheet(unknown, unknown_headers, unknown_rows)

    well_sheet = wb.create_sheet("Well Results")
    well_headers = ["Well", "Row", "Column", "Source cell", "Original reading", "Raw OD", "Measurement status", "Role", "Include", "Exclusion reason", "Standard curve ID", "Nominal concentration", "Sample ID", "Replicate group", "Dilution factor", "Corrected OD", "Measured concentration", "Dilution-adjusted concentration", "Standard predicted OD", "Standard residual OD", "Standard recovery %", "Result status", "Flags"]
    keys = ("well", "row", "column", "source_cell", "raw_value", "raw_od", "measurement_status", "role", "include", "exclusion_reason", "curve_id", "nominal_concentration", "sample_id", "replicate_group", "dilution_factor", "corrected_od", "well_concentration", "well_adjusted_concentration", "standard_predicted_od", "standard_residual_od", "standard_recovery_pct", "result_status", "flags")
    _sheet(well_sheet, well_headers, [[r.get(k) for k in keys] for r in result["well_results"]])

    standards = wb.create_sheet("Standards & QC")
    headers = ["Standard curve ID", "Nominal concentration", "Assigned n", "Usable n", "Wells", "Measurement statuses", "Original readings", "Mean corrected OD", "OD SD", "OD CV %", "Back-calculated concentration by well", "Recovery % by well", "Residual OD by well", "Flags"]
    keys = ("curve_id", "nominal_concentration", "n_assigned", "n_usable", "wells", "statuses", "readings", "mean_corrected_od", "sd_corrected_od", "cv_pct", "back_calculated", "recovery_pct", "residual_od", "flags")
    _sheet(standards, headers, [[r[k] for k in keys] for r in result["standards_qc"]])

    fits = wb.create_sheet("Standard Curves")
    fit_headers = ["Standard curve ID", "Units", "Model", "Status", "Fit status", "QC status", "Usable levels", "Usable wells", "Direction", "Equation", "Fitted equation", "Bottom", "Top", "EC50", "Slope", "Range low", "Range high", "Response low", "Response high", "RMSE OD", "R squared (supplementary)", "Jacobian condition", "Global blank mean OD", "Flags"]
    fit_rows = []
    for c in result["curves"]:
        p = c["parameters"] or {}
        fit_rows.append([c["curve_id"], c["units"], c["model_type"], c["status"], c["fit_status"], c["qc_status"], c["n_levels"], c["n_wells"], c["direction"], c["equation"], c["fitted_equation"], p.get("bottom"), p.get("top"), p.get("ec50"), p.get("slope"), c["range_low"], c["range_high"], c["response_low"], c["response_high"], c["rmse_od"], c["r_squared"], c.get("jacobian_condition"), c["blank_mean_od"], c["flags"]])
    _sheet(fits, fit_headers, fit_rows)
    plot_row = max(5, len(fit_rows) + 4)
    for curve in result["curves"]:
        image = Image(BytesIO(curve_plot_bytes(curve, result["well_results"], "png")))
        image.width, image.height = 750, 440
        fits.add_image(image, f"A{plot_row}")
        plot_row += 25

    setting_sheet = wb.create_sheet("Settings")
    setting_rows = [["Software version", __version__], ["Layout JSON", "Download separately from application; annotations below are an audit copy."]]
    setting_rows += [[k, v] for k, v in result["settings"].items()]
    setting_rows.append(["primary_method_meaning", method_meaning])
    setting_rows += [["Standard curve " + cid + " units", c["units"]] for cid, c in layout.get("curves", {}).items()]
    setting_rows += [["Annotation " + well, str(annotation)] for well, annotation in layout.get("annotations", {}).items()]
    _sheet(setting_sheet, ["Setting", "Value"], setting_rows, [32, 105])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
